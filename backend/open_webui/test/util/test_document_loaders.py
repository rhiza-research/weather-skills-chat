import importlib.util
import socket
import tempfile
import unittest
import zipfile
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from fpdf import FPDF
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from open_webui.retrieval.loaders.main import Loader

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EPUB_TYPE = "application/epub+zip"


@contextmanager
def no_network():
    """Fail any connection or DNS lookup, and record it so a swallowed error still shows."""
    calls: list[tuple[str, tuple]] = []

    def blocked(name, error):
        def _blocked(*args, **kwargs):
            calls.append((name, args))
            raise error("network blocked in test")

        return _blocked

    with (
        patch.object(socket.socket, "connect", blocked("connect", OSError)),
        patch.object(socket.socket, "connect_ex", blocked("connect_ex", OSError)),
        patch.object(socket, "getaddrinfo", blocked("getaddrinfo", socket.gaierror)),
    ):
        yield calls


def write_epub(path: Path, title: str, body: str) -> None:
    container = (
        '<?xml version="1.0"?>'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>'
    )
    opf = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
        f'<dc:identifier id="id">test-epub</dc:identifier><dc:title>{title}</dc:title>'
        "<dc:language>en</dc:language></metadata>"
        '<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" '
        'properties="nav"/><item id="ch1" href="chapter1.xhtml" '
        'media-type="application/xhtml+xml"/></manifest>'
        '<spine><itemref idref="ch1"/></spine></package>'
    )
    nav = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">'
        '<head><title>Nav</title></head><body><nav epub:type="toc"><ol>'
        '<li><a href="chapter1.xhtml">Chapter 1</a></li></ol></nav></body></html>'
    )
    chapter = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter 1</title></head>'
        f"<body><h1>{title}</h1><p>{body}</p></body></html>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("META-INF/container.xml", container, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/nav.xhtml", nav, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/chapter1.xhtml", chapter, compress_type=zipfile.ZIP_DEFLATED)


def write_docx(path: Path, text: str) -> None:
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/'
        'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("word/document.xml", document)


def write_text_pdf(path: Path, text: str) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.cell(text=text)
    pdf.output(str(path))


def write_image_pdf(path: Path, text: str) -> None:
    """A PDF whose only content is a JPEG of rendered text, so its text comes from OCR."""
    image = Image.new("RGB", (1000, 240), "white")
    ImageDraw.Draw(image).text(
        (60, 60), text, fill="black", font=ImageFont.load_default(size=96)
    )
    jpeg = path.with_suffix(".jpg")
    image.save(jpeg, format="JPEG", quality=95)
    pdf = FPDF()
    pdf.add_page()
    pdf.image(str(jpeg), x=10, y=20, w=190)
    pdf.output(str(path))


class DocumentLoaderTest(unittest.TestCase):
    def _load(self, path: Path, content_type: str, **loader_kwargs) -> str:
        with no_network() as calls:
            docs = Loader(**loader_kwargs).load(path.name, content_type, str(path))
        self.assertEqual(calls, [])
        return "\n".join(doc.page_content for doc in docs)

    def test_rst(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "notes.rst"
            path.write_text(
                "Rainfall Notes\n==============\n\nKenya rainfall was above normal.\n",
                encoding="utf-8",
            )
            text = self._load(path, "text/x-rst")
        self.assertIn("Rainfall Notes", text)
        self.assertIn("Kenya rainfall was above normal.", text)

    def test_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "stations.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["station", "rainfall_mm"])
            ws.append(["Nairobi", 42])
            wb.save(path)
            text = self._load(path, XLSX_TYPE)
        self.assertIn("Nairobi", text)
        self.assertIn("rainfall_mm", text)

    def test_pptx(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "briefing.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "Weekly Briefing"
            slide.placeholders[1].text = "Heavy rain expected in Mombasa."
            prs.save(path)
            text = self._load(path, PPTX_TYPE)
        self.assertIn("Weekly Briefing", text)
        self.assertIn("Heavy rain expected in Mombasa.", text)

    def test_xml(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.xml"
            path.write_text(
                '<?xml version="1.0"?><report><station>Kisumu</station>'
                "<summary>Light showers over the lake.</summary></report>",
                encoding="utf-8",
            )
            text = self._load(path, "application/xml")
        self.assertIn("Kisumu", text)
        self.assertIn("Light showers over the lake.", text)

    def test_epub(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "guide.epub"
            write_epub(path, "Field Guide", "Rain gauges are read at nine each morning.")
            text = self._load(path, EPUB_TYPE)
        self.assertIn("Field Guide", text)
        self.assertIn("Rain gauges are read at nine each morning.", text)

    def test_cp1252_text_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "notes.txt"
            path.write_bytes("Café au lait — résumé of the week.".encode("cp1252"))
            text = self._load(path, "text/plain")
        self.assertIn("Café au lait", text)
        self.assertIn("résumé", text)
        # cp1252 0x97 is an em dash; Latin-1 would decode it as a control character.
        self.assertIn("—", text)

    def test_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "outlook.pdf"
            write_text_pdf(path, "Rainfall outlook for Nakuru")
            text = self._load(path, "application/pdf")
        self.assertIn("Rainfall outlook for Nakuru", text)

    def test_pdf_image_ocr(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "scan.pdf"
            write_image_pdf(path, "RAINFALL")
            text = self._load(path, "application/pdf", PDF_EXTRACT_IMAGES=True)
        self.assertIn("RAINFALL", text)

    def test_csv(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "stations.csv"
            path.write_text("station,rainfall_mm\nNairobi,42\n", encoding="utf-8")
            text = self._load(path, "text/csv")
        self.assertIn("station: Nairobi", text)
        self.assertIn("rainfall_mm: 42", text)

    def test_html(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bulletin.html"
            path.write_text(
                "<html><head><title>Coast Bulletin</title></head>"
                "<body><p>Strong winds along the coast.</p></body></html>",
                encoding="utf-8",
            )
            text = self._load(path, "text/html")
        self.assertIn("Coast Bulletin", text)
        self.assertIn("Strong winds along the coast.", text)

    def test_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.docx"
            write_docx(path, "Station report for Eldoret.")
            text = self._load(path, DOCX_TYPE)
        self.assertIn("Station report for Eldoret.", text)


class SpacyModelTest(unittest.TestCase):
    def test_unstructured_tokenizer_loads_model_offline(self):
        import spacy.cli
        from unstructured.nlp import tokenize

        self.assertIsNotNone(
            importlib.util.find_spec("en_core_web_sm"),
            "en_core_web_sm is not installed in the environment",
        )

        # unstructured caches the loaded model; clear it so the load below runs offline.
        for name in ("_get_nlp", "_process_cached", "_tokenize_for_cache"):
            cached = getattr(tokenize, name, None)
            self.assertIsNotNone(
                cached,
                f"unstructured.nlp.tokenize.{name} no longer exists; unstructured changed "
                "how it caches the spaCy model, so this test must clear the new cache",
            )
            cached.cache_clear()
        self.assertTrue(
            hasattr(tokenize, "_install_spacy_model"),
            "unstructured.nlp.tokenize._install_spacy_model no longer exists; unstructured "
            "changed how it installs a missing spaCy model, so this test must block the new path",
        )

        def refuse_download(*args, **kwargs):
            raise AssertionError("tried to download the spaCy model instead of loading it")

        with (
            patch.object(spacy.cli, "download", refuse_download),
            patch.object(tokenize, "_install_spacy_model", refuse_download),
            no_network() as calls,
        ):
            sentences = tokenize.sent_tokenize(
                "Rain fell in Nairobi overnight. Floods followed in Mombasa."
            )
        self.assertEqual(calls, [])
        self.assertEqual(
            sentences,
            ["Rain fell in Nairobi overnight.", "Floods followed in Mombasa."],
        )


if __name__ == "__main__":
    unittest.main()
